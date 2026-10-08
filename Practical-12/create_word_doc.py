"""
Word Document Generator for BDA Practical 12
Matching Practical 8 Lab Manual Structure with Real Terminal Screenshots
Author: Sil Shah (24DIT066)
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
IMG_DIR = os.path.join(BASE_DIR, "images")

def add_screenshot(doc, img_path, width_inch=5.8):
    if os.path.exists(img_path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(8)
        run = p.add_run()
        run.add_picture(img_path, width=Inches(width_inch))

def generate_doc():
    doc = docx.Document()
    
    # Page Margins & Header/Footer
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)
        
        # Header: CSUE301 : Big Data Analytics (Left) | 24DIT066 (Right)
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        
        hrun1 = hp.add_run("CSUE301 : Big Data Analytics" + " " * 85)
        hrun1.font.name = "Calibri"
        hrun1.font.size = Pt(9.5)
        hrun1.font.color.rgb = RGBColor(60, 60, 60)
        
        hrun2 = hp.add_run("24DIT066")
        hrun2.font.bold = True
        hrun2.font.name = "Calibri"
        hrun2.font.size = Pt(9.5)
        hrun2.font.color.rgb = RGBColor(0, 0, 0)
        
        # Footer: Page number
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        frun = fp.add_run("")
        fldSimple = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
        fp._p.append(fldSimple)

    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(30, 30, 30)

    # ------------------ PAGE 1: COVER PAGE ------------------
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(60)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("CHAROTAR UNIVERSITY OF SCIENCE & TECHNOLOGY\n")
    r.font.bold = True
    r.font.size = Pt(16)
    r.font.color.rgb = RGBColor(0, 51, 102)

    r = p.add_run("Faculty of Technology & Engineering\n")
    r.font.bold = True
    r.font.size = Pt(14)
    
    r = p.add_run("Devang Patel Institute of Advance Technology and Research (DEPSTAR)\n\n")
    r.font.bold = True
    r.font.size = Pt(13)

    r = p.add_run("CSUE301 – : Big Data Analytics\n")
    r.font.bold = True
    r.font.size = Pt(14)

    r = p.add_run("Practical File – ODD Semester 2026-27\n\n")
    r.font.bold = True
    r.font.size = Pt(13)

    r = p.add_run("StudentID: 24DIT066\n")
    r.font.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = RGBColor(0, 51, 102)

    r = p.add_run("Student Name: Sil Shah\n")
    r.font.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = RGBColor(0, 51, 102)

    doc.add_page_break()

    # ------------------ PAGE 2: PRACTICAL 12 ------------------
    h1 = doc.add_paragraph()
    h1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h1.paragraph_format.space_before = Pt(8)
    h1.paragraph_format.space_after = Pt(12)
    r = h1.add_run("Practical 12")
    r.font.bold = True
    r.font.size = Pt(18)
    r.font.color.rgb = RGBColor(0, 51, 102)

    # Problem Definition
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("1. Problem Definition\n")
    r.font.bold = True
    r.font.size = Pt(12)
    doc.add_paragraph(
        "You are working as a Data Visualization Analyst at a retail analytics company. Management wants to "
        "monitor sales performance across products, regions, and customer segments through an interactive dashboard. "
        "Static reports are insufficient for decision-making. Your task is to design an interactive dashboard that allows "
        "management to explore sales data and identify key business trends."
    )

    # PART A
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("PART A — Practical Execution\n")
    r.font.bold = True
    r.font.size = Pt(12)

    # Step 1
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("Step 1: Creating Folder and files")
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 51, 102)
    add_screenshot(doc, os.path.join(IMG_DIR, "step1_folder.png"), 3.2)

    # Step 2
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("Step 2: Loading & Validating Superstore Sales Dataset")
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 51, 102)
    add_screenshot(doc, os.path.join(IMG_DIR, "step2_load_validate.png"), 5.8)

    # Step 3
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("Step 3: Display Dataset Schema & Preprocessing")
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 51, 102)
    add_screenshot(doc, os.path.join(IMG_DIR, "step3_schema_dates.png"), 5.8)

    # Step 4
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("Step 4: Executive KPI Calculation")
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 51, 102)
    add_screenshot(doc, os.path.join(IMG_DIR, "step4_kpi_metrics.png"), 5.8)

    # Step 5
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("Step 5: Category-Level Sales & Profit Performance")
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 51, 102)
    add_screenshot(doc, os.path.join(IMG_DIR, "step5_category_perf.png"), 5.8)

    # Step 6
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("Step 6: Sub-Category Profitability Drill-Down")
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 51, 102)
    add_screenshot(doc, os.path.join(IMG_DIR, "step6_subcat_drilldown.png"), 5.8)

    # Step 7
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("Step 7: Regional Sales & Profit Distribution")
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 51, 102)
    add_screenshot(doc, os.path.join(IMG_DIR, "step7_regional_dist.png"), 5.8)

    # Step 8
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("Step 8: Customer Segment Dynamics & Basket Size")
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 51, 102)
    add_screenshot(doc, os.path.join(IMG_DIR, "step8_segment_breakdown.png"), 5.8)

    # Step 9
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("Step 9: Quarterly Performance & QoQ Growth Analysis")
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 51, 102)
    add_screenshot(doc, os.path.join(IMG_DIR, "step9_quarterly_timeline.png"), 5.8)

    # Step 10
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("Step 10: Interactive Dashboard Application Server Launch")
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 51, 102)
    add_screenshot(doc, os.path.join(IMG_DIR, "step10_streamlit_launch.png"), 5.8)

    # Step 11
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("Step 11: Executive KPI Summary Dashboard Visualization")
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 51, 102)
    add_screenshot(doc, os.path.join(BASE_DIR, "kpi_summary_dashboard.png"), 6.2)

    # Step 12
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("Step 12: Category Drill-Down Profitability Visualization")
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 51, 102)
    add_screenshot(doc, os.path.join(BASE_DIR, "category_drilldown_analysis.png"), 6.0)

    # Step 13
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("Step 13: Regional & Top State Performance Visualization")
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 51, 102)
    add_screenshot(doc, os.path.join(BASE_DIR, "regional_performance_map.png"), 5.6)

    # Step 14
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("Step 14: Quarterly Business Performance Trend Visualization")
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 51, 102)
    add_screenshot(doc, os.path.join(BASE_DIR, "monthly_quarterly_trend.png"), 5.6)

    # Step 15
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("Step 15: Customer Segmentation Analysis Visualization")
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 51, 102)
    add_screenshot(doc, os.path.join(BASE_DIR, "customer_segmentation_analysis.png"), 5.6)

    # Step 16
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("Step 16: Generate Final Sales Dashboard Summary Report")
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 51, 102)
    add_screenshot(doc, os.path.join(IMG_DIR, "step16_report_preview.png"), 5.8)

    # Section 3: Key Questions
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("3. Key Questions / Analysis\n")
    r.font.bold = True
    r.font.size = Pt(12)

    # Q1
    p = doc.add_paragraph()
    r = p.add_run("Q1. What is the role of interactive dashboards in business analytics?\n")
    r.font.bold = True
    doc.add_paragraph(
        "Interactive dashboards translate complex underlying datasets into visual metrics and KPIs. "
        "They allow decision-makers to track performance in real-time, test business hypotheses, spot anomalies, "
        "and explore data across multiple dimensions without needing to write custom queries or examine raw tabular spreadsheets."
    )

    # Q2
    p = doc.add_paragraph()
    r = p.add_run("Q2. How do filters improve visual exploration?\n")
    r.font.bold = True
    doc.add_paragraph(
        "Filters allow users to isolate specific segments of data (such as a single region, category, or year). "
        "By removing noise and focusing on relevant subsets, cross-filtering reveals relationships and patterns "
        "between different business dimensions that might otherwise remain hidden."
    )

    # Q3
    p = doc.add_paragraph()
    r = p.add_run("Q3. What are KPI indicators?\n")
    r.font.bold = True
    doc.add_paragraph(
        "Key Performance Indicators (KPIs) are quantifiable metrics used to track and evaluate organizational performance "
        "against strategic goals. Examples in retail analytics include Total Sales, Net Profit, Profit Margin %, "
        "Average Order Value (AOV), and Customer Retention Rate."
    )

    # Q4
    p = doc.add_paragraph()
    r = p.add_run("Q4. Why are dashboards preferred over static reports?\n")
    r.font.bold = True
    doc.add_paragraph(
        "Dashboards are preferred over static reports because they support interactive drill-downs, dynamic filtering, "
        "and instant cross-filtering. Static reports only show fixed snapshots that cannot be queried further, "
        "whereas interactive dashboards enable users to discover root causes behind trends on demand."
    )

    # Q5
    p = doc.add_paragraph()
    r = p.add_run("Q5. How does user-driven exploration support decision-making?\n")
    r.font.bold = True
    doc.add_paragraph(
        "User-driven exploration allows managers to ask follow-up questions immediately as they spot trends. "
        "For example, seeing a low profit margin in Furniture allows the user to drill down into sub-categories to discover "
        "that heavy discounts on Tables and Bookcases are causing the loss, enabling targeted strategic interventions."
    )

    # Section 4: Conclusion
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("4. Conclusion\n")
    r.font.bold = True
    r.font.size = Pt(12)
    doc.add_paragraph(
        "The practical successfully demonstrated the creation and deployment of an interactive Customer Sales Dashboard using Business Intelligence and interactive analytics.\n"
        "A multi-year Superstore Sales dataset of 5,000 records was cleaned and validated. Core KPIs were formulated, yielding $8,885,748.66 in total sales and $683,907.28 in profit (7.70% margin). "
        "Multi-dimensional drill-downs identified Technology as the strongest profit driver (14.23% margin, led by Copiers and Phones) and Furniture as a key loss center (-8.86% margin) due to excessive discounting.\n"
        "Regional analysis confirmed the West region as the highest contributor to net profit ($195.0k), while customer segmentation showed Home Office producing the highest profit margin (9.35%). "
        "An interactive web dashboard application was deployed using Streamlit and Plotly, providing live filtering, interactive visuals, and drill-through table exploration."
    )

    # Section 5: Github Link
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("5. Github Link\n")
    r.font.bold = True
    r.font.size = Pt(12)
    
    p = doc.add_paragraph()
    r1 = p.add_run("Github Link: _ ")
    r2 = p.add_run("24DIT066-tech/Practical-12 at main · sil-sys/24DIT066-tech")
    r2.font.color.rgb = RGBColor(0, 102, 204)
    r2.font.underline = True

    output_path = os.path.join(BASE_DIR, "Practical_12_Customer_Sales_Dashboard.docx")
    doc.save(output_path)
    print(f"Generated Practical 12 Word Document at: {output_path}")

if __name__ == "__main__":
    generate_doc()
